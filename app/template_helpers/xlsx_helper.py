"""
xlsx_helper.py
---------------

OpenPyXL-based helper utilities that surface placeholders from XLSX templates.
Consultants can use the resulting JSON report to craft precise prompts or to
drive automated filling inside Code Interpreter.
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from openpyxl import load_workbook  # type: ignore

PLACEHOLDER_PATTERNS = (
    re.compile(r"\[\[.+?\]\]"),
    re.compile(r"\{\{.+?\}\}"),
    re.compile(r"\bInsert\b", re.IGNORECASE),
)


@dataclass
class CellPlaceholder:
    sheet: str
    cell: str
    value: str


def _matches(text: str | None) -> bool:
    if not text:
        return False
    return any(pattern.search(str(text)) for pattern in PLACEHOLDER_PATTERNS)


def _iter_cells(path: Path) -> Iterable[Tuple[str, str, str]]:
    workbook = load_workbook(path, data_only=True)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if _matches(value):
                    yield sheet.title, cell.coordinate, str(value).strip()


def scan_xlsx_placeholders(path: Path | str) -> List[CellPlaceholder]:
    template = Path(path)
    if not template.exists():
        raise FileNotFoundError(f"XLSX template not found: {template}")
    return [
        CellPlaceholder(sheet=sheet, cell=cell, value=value)
        for sheet, cell, value in _iter_cells(template)
    ]


def export_xlsx_report(path: Path | str, output_json: Path | str | None = None) -> str:
    placeholders = scan_xlsx_placeholders(path)
    payload = {
        "template": str(path),
        "placeholder_count": len(placeholders),
        "placeholders": [asdict(p) for p in placeholders],
    }
    json_blob = json.dumps(payload, indent=2)
    if output_json:
        Path(output_json).write_text(json_blob, encoding="utf-8")
    return json_blob


def _cli(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Inspect an XLSX template and list placeholder cells."
    )
    parser.add_argument("template", type=Path, help="Path to the XLSX file.")
    parser.add_argument("--out", type=Path, help="Optional JSON output file.")
    args = parser.parse_args(argv)
    json_blob = export_xlsx_report(args.template, args.out)
    print(json_blob)
    return 0


if __name__ == "__main__":
    raise SystemExit(_cli())

