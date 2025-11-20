"""
docx_helper.py
---------------

Utility routines that make it easier to inspect Limiting Competition DOCX templates.
They scan all paragraphs and table cells, surface likely placeholders, and support a
CLI so prompts/system messages can embed precise placeholder labels instead of
generic “Insert ...” text.
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Iterable, List, Sequence

from docx import Document  # type: ignore

PLACEHOLDER_PATTERNS: Sequence[re.Pattern[str]] = (
    re.compile(r"\[\[.+?\]\]"),
    re.compile(r"\bInsert\b", re.IGNORECASE),
)


@dataclass
class Placeholder:
    location: str
    text: str
    context: str


def _matches_placeholder(text: str) -> bool:
    normalized = text.strip()
    if not normalized:
        return False
    return any(pattern.search(normalized) for pattern in PLACEHOLDER_PATTERNS)


def _scan_paragraphs(doc: Document) -> Iterable[Placeholder]:
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text or ""
        if _matches_placeholder(text):
            yield Placeholder(
                location=f"paragraph[{idx}]",
                text=text.strip(),
                context=paragraph.style.name if paragraph.style else "",
            )


def _scan_tables(doc: Document) -> Iterable[Placeholder]:
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                text = cell.text or ""
                if _matches_placeholder(text):
                    yield Placeholder(
                        location=f"table[{table_idx}].row[{row_idx}].col[{col_idx}]",
                        text=text.strip(),
                        context=cell.paragraphs[0].style.name if cell.paragraphs else "",
                    )


def scan_docx_placeholders(path: Path | str) -> List[Placeholder]:
    """Return all placeholders detected inside the provided DOCX template."""
    template_path = Path(path)
    if not template_path.exists():
        raise FileNotFoundError(f"DOCX template not found: {template_path}")
    document = Document(template_path)
    hits = list(_scan_paragraphs(document))
    hits.extend(_scan_tables(document))
    return hits


def export_docx_report(path: Path | str, output_json: Path | str | None = None) -> str:
    """
    Generate a JSON report highlighting every placeholder hit.
    Returns the JSON string and optionally writes it to ``output_json``.
    """
    placeholders = scan_docx_placeholders(path)
    payload = {
        "template": str(path),
        "placeholder_count": len(placeholders),
        "placeholders": [asdict(p) for p in placeholders],
    }
    json_blob = json.dumps(payload, indent=2)
    if output_json:
        Path(output_json).write_text(json_blob, encoding="utf-8")
    return json_blob


def _cli(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Inspect a DOCX template and list placeholder candidates."
    )
    parser.add_argument("template", type=Path, help="Path to the DOCX template.")
    parser.add_argument(
        "--out",
        type=Path,
        default=None,
        help="Optional JSON file path for saving the report.",
    )
    args = parser.parse_args(argv)
    json_blob = export_docx_report(args.template, args.out)
    print(json_blob)
    return 0


if __name__ == "__main__":
    raise SystemExit(_cli())

