"""
Fuzzy anchor mapping helpers.

Given a template (DOCX/XLSX/PDF) and user-defined targets (e.g., "Description of Action"),
suggest the best matching anchors so humans can confirm/edit before Code Interpreter runs.
"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass, asdict
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, Iterable, List, Sequence

from docx import Document  # type: ignore
from openpyxl import load_workbook  # type: ignore

from .pdf_helper import extract_pdf_text


@dataclass
class Chunk:
    source: str
    location: str
    text: str


def _collect_docx_chunks(path: Path) -> List[Chunk]:
    doc = Document(path)
    chunks: List[Chunk] = []
    for idx, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if text:
            chunks.append(Chunk("paragraph", f"paragraph[{idx}]", text))
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                text = (cell.text or "").strip()
                if text:
                    chunks.append(
                        Chunk(
                            "table",
                            f"table[{table_idx}].row[{row_idx}].col[{col_idx}]",
                            text,
                        )
                    )
    return chunks


def _collect_xlsx_chunks(path: Path) -> List[Chunk]:
    wb = load_workbook(path, data_only=True)
    chunks: List[Chunk] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                text = str(value).strip()
                if not text:
                    continue
                chunks.append(
                    Chunk(
                        "sheet",
                        f"{sheet.title}!{cell.coordinate}",
                        text,
                    )
                )
    return chunks


def _collect_pdf_chunks(path: Path) -> List[Chunk]:
    pages = extract_pdf_text(path)
    chunks: List[Chunk] = []
    for page in pages:
        for line_idx, line in enumerate(page.text.splitlines()):
            text = line.strip()
            if text:
                chunks.append(Chunk("pdf", f"page[{page.page}].line[{line_idx}]", text))
    return chunks


def _collect_chunks(path: Path) -> List[Chunk]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _collect_docx_chunks(path)
    if suffix == ".xlsx":
        return _collect_xlsx_chunks(path)
    if suffix == ".pdf":
        return _collect_pdf_chunks(path)
    raise ValueError(f"Unsupported template type: {suffix}")


def _score(target: str, text: str) -> float:
    target_norm = target.strip().lower()
    text_norm = text.strip().lower()
    if not target_norm or not text_norm:
        return 0.0
    return SequenceMatcher(None, target_norm, text_norm).ratio()


def find_fuzzy_matches(
    template_path: Path | str, targets: Sequence[str], top_k: int = 3
) -> Dict[str, List[Dict[str, float]]]:
    """Return top matches (with scores) for each target string."""
    path = Path(template_path)
    chunks = _collect_chunks(path)
    results: Dict[str, List[Dict[str, float]]] = {}
    for target in targets:
        scored = []
        for chunk in chunks:
            score = _score(target, chunk.text)
            if score <= 0:
                continue
            scored.append({"score": score, "location": chunk.location, "text": chunk.text})
        scored.sort(key=lambda x: x["score"], reverse=True)
        results[target] = scored[:top_k]
    return results


def propose_anchor_map(
    template_path: Path | str,
    targets: Sequence[str],
    top_k: int = 3,
    confidence_threshold: float = 0.82,
) -> Dict[str, Dict[str, object]]:
    """Generate a reviewable map for humans: best match + alternates + confidence flag."""
    matches = find_fuzzy_matches(template_path, targets, top_k=top_k)
    proposal: Dict[str, Dict[str, object]] = {}
    for target, candidates in matches.items():
        top = candidates[0] if candidates else None
        proposal[target] = {
            "auto_selection": top,
            "candidates": candidates,
            "requires_review": (not top) or (top["score"] < confidence_threshold),
        }
    return proposal


def _cli(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Suggest fuzzy anchor matches for a template."
    )
    parser.add_argument("template", type=Path)
    parser.add_argument(
        "--targets",
        type=Path,
        required=True,
        help="Path to a JSON file of target strings (array).",
    )
    parser.add_argument("--out", type=Path, help="Optional output JSON path.")
    parser.add_argument("--top", type=int, default=3)
    parser.add_argument("--threshold", type=float, default=0.82)
    args = parser.parse_args(argv)
    targets = json.loads(args.targets.read_text())
    proposal = propose_anchor_map(
        args.template,
        targets,
        top_k=args.top,
        confidence_threshold=args.threshold,
    )
    blob = json.dumps(proposal, indent=2)
    if args.out:
        args.out.write_text(blob, encoding="utf-8")
    print(blob)
    return 0


if __name__ == "__main__":
    raise SystemExit(_cli())

